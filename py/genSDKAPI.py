# !/usr/bin/python
# encoding:utf-8
import json
import re
from docx import Document

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt

from utils.fileUtils import findFilePathNameByKeyWord, parse_cmake_versions




def set_cell_background(cell, colorstr):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{color_value}"/>'.format(nsdecls('w'), color_value=colorstr))
    cell._tc.get_or_add_tcPr().append(shading_elm)




def readAPI(APIpathName, keyword, funkey):
    """
    收集 API 声明行：同时含 keyword（如 OXISCAN_API）和 funkey（如 OXiScan_）。

    按「函数名」去重、保持首次出现的顺序。头文件里同名函数可能被整块重复声明
    （实测 base/export/OXiScan_Api.h 里 OXiScan_getChipID 连注释带声明出现了两遍），
    不去重的话文档里会出现两条同名条目，且**后面所有小节的编号整体错位**。
    """
    liner = []
    seen = set()
    with open(file=APIpathName, encoding="utf-8") as f:
        for line in f:
            if keyword not in line or funkey not in line:
                continue
            funName = extract_function_name(line, verbose=False)
            # 函数名解析不出来时退化成整行比对，避免把不同的行误合并
            key = funName if funName else line.strip()
            if key in seen:
                print("跳过重复声明:", key)
                continue
            seen.add(key)
            liner.append(line)
    return liner


def extract_function_name(funStr, verbose=True):

    pattern = re.compile(r'\b(?:[a-zA-Z_][a-zA-Z0-9_]*::)?(?:template\s*<.*?>\s+)?([a-zA-Z_][a-zA-Z0-9_]*)\s*\([^)]*\)\s*(?:const)?\s*(?:throw\s*\([^)]*\))?\s*;\s*|\b(?:[a-zA-Z_][a-zA-Z0-9_]*::)?(?:template\s*<.*?>\s+)?([a-zA-Z_][a-zA-Z0-9_]*)\s*\([^)]*\)\s*(?:const)?\s*(?:\{|\s*throw\s*\([^)]*\)\s*;\s*)', re.MULTILINE)

    match = pattern.search(funStr)
    if match:
        function_name = next((group for group in match.groups() if group), None)
        if verbose:
            print("第一个匹配到的函数名是:", function_name)
        return function_name
    else:
        if verbose:
            print("没有找到匹配项")

    return ""


def getFuntionString(funStr, keyword):
    return funStr.replace(keyword +" ", "").strip()



def extract_brief_description(funName, APIPathName):

    f = open(file=APIPathName, encoding="utf-8")
    liner = []
    havaBreak = True
    for line in f:
        
        liner.append(line.replace('\n', '').replace('\r', '').replace('\r\n', ''))
        if "brief" in line:
            if havaBreak:
                liner.clear()
                liner.append(line.replace('\n', '').replace('\r', '').replace('\r\n', ''))
            else:
                havaBreak = True
           
        if funName.lower() in line.lower():
            break
                
    return liner


# ---------------------------------------------------------------
# 目录（TOC）
#
# 参考文档 D:\GSerSDKLib\doc\Fingerprint Scanner SDK (ISO)v 2.0.0.31.docx 的第 2 页是目录：
#   CONTEXT 标题 + TOC 域（\o "1-3"），逐条列出 1/2/3 章，以及 2 章下的 4 个二级小节
#   （Supported Fingerprint Sensor / SDK Features / Support for / Supported Communication
#   Interface）和 3 章下的每个 API 小节。
# 本脚本生成的文档原先整块缺目录，这里补齐。
#
# 构造方式与 Word 生成的目录完全一致：
#   正文每个标题带一个书签（_TocXXXXXX） → 目录条目是超链接 + PAGEREF 域指向该书签。
# 条目文本与页码是「域缓存值」，打开文档即可见；在 Word 里 Ctrl+A → F9 可刷新为真实页码。
# 脚本本身无法分页，页码按参考文档 2.0.0.31 的排版规律预估（封面1、目录2、1章3、2章4、
# 3章及第 1 个 API 在 5、之后每 2 个 API 翻一页）。

_TOC_ENTRIES = []
_BOOKMARK_ID = [1000]


def _xml_escape(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def add_toc_heading(doc, text, level, page):
    """
    等价于 doc.add_heading(text, level)，但额外：
      1) 给标题段落加一对书签（目录的 PAGEREF 要引用它）；
      2) 把 (层级, 文本, 页码) 登记进 _TOC_ENTRIES，供最后统一生成目录页。
    """
    heading = doc.add_heading(text, level)
    anchor = "_Toc%d" % (900000 + len(_TOC_ENTRIES))
    bid = _BOOKMARK_ID[0]
    _BOOKMARK_ID[0] += 1
    p = heading._p
    ppr = p.find(qn("w:pPr"))
    start = parse_xml('<w:bookmarkStart %s w:id="%d" w:name="%s"/>' % (nsdecls("w"), bid, anchor))
    end = parse_xml('<w:bookmarkEnd %s w:id="%d"/>' % (nsdecls("w"), bid))
    if ppr is not None:
        ppr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)
    _TOC_ENTRIES.append({"level": level, "text": text, "anchor": anchor, "page": page})
    return heading


def _toc_entry_xml(entry, is_first):
    """目录里的一条：右对齐点线制表位 + 超链接 + 页码（PAGEREF 域）。"""
    indent = "" if entry["level"] <= 1 else '<w:ind w:left="440"/>'
    rpr = '<w:rPr><w:noProof/></w:rPr>'
    rpr_hidden = '<w:rPr><w:noProof/><w:webHidden/></w:rPr>'
    # 第一条还要带上 TOC 域的开头（begin / 域代码 / separate）
    head = ""
    if is_first:
        head = (
            '<w:r>' + rpr + '<w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r>' + rpr + '<w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
            '<w:r>' + rpr + '<w:fldChar w:fldCharType="separate"/></w:r>'
        )
    return (
        '<w:p><w:pPr>'
        '<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="8630"/></w:tabs>'
        + indent +
        '</w:pPr>'
        + head +
        '<w:hyperlink w:anchor="%s" w:history="1">' % entry["anchor"]
        + '<w:r>' + rpr + '<w:t xml:space="preserve">' + _xml_escape(entry["text"]) + '</w:t></w:r>'
        + '<w:r>' + rpr_hidden + '<w:tab/></w:r>'
        + '<w:r>' + rpr_hidden + '<w:fldChar w:fldCharType="begin"/></w:r>'
        + '<w:r>' + rpr_hidden + '<w:instrText xml:space="preserve"> PAGEREF %s \\h </w:instrText></w:r>' % entry["anchor"]
        + '<w:r>' + rpr_hidden + '<w:fldChar w:fldCharType="separate"/></w:r>'
        + '<w:r>' + rpr_hidden + '<w:t>%d</w:t></w:r>' % entry["page"]
        + '<w:r>' + rpr_hidden + '<w:fldChar w:fldCharType="end"/></w:r>'
        + '</w:hyperlink>'
        '</w:p>'
    )


def build_toc_xml(entries):
    """整块目录：内容控件（sdt，与 Word 的「目录」部件同类）+ CONTEXT 标题 + 各条目 + 域收尾。"""
    parts = ['<w:sdt %s>' % nsdecls("w")]
    parts.append(
        '<w:sdtPr><w:id w:val="18072631"/>'
        '<w:docPartObj><w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/></w:docPartObj>'
        '</w:sdtPr>'
        '<w:sdtContent>'
    )
    parts.append('<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>CONTEXT</w:t></w:r></w:p>')
    for index, entry in enumerate(entries):
        parts.append(_toc_entry_xml(entry, index == 0))
    # 收尾：关闭 TOC 域
    parts.append(
        '<w:p><w:pPr><w:rPr><w:noProof/></w:rPr></w:pPr>'
        '<w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="end"/></w:r></w:p>'
    )
    parts.append('</w:sdtContent></w:sdt>')
    return "".join(parts)


def insert_toc(doc, first_body_paragraph, entries):
    """把目录页插到正文第一个标题之前（即封面之后），并在目录后面补一个分页符。"""
    if not entries:
        return
    node = parse_xml(build_toc_xml(entries))
    first_body_paragraph.addprevious(node)
    page_break = parse_xml('<w:p %s><w:r><w:br w:type="page"/></w:r></w:p>' % nsdecls("w"))
    node.addnext(page_break)


def write_revision_record(doc):
    heading = add_toc_heading(doc, "1. Record of Revision", 1, page=3)
    table = doc.add_table(rows=1, cols=4, style="Table Grid")
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Rev'
    hdr_cells[1].text = 'Issue Date'
    hdr_cells[2].text = 'Description'
    hdr_cells[3].text = 'Editor'

    for cell in table.rows[0].cells:
        set_cell_background(cell, 'c0c0c0')
        cell.height = Inches(1.5)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for index in range(4):
        row_cells = table.add_row().cells
        table.rows[index+1].height = Inches(0.3)
        row_cells[0].width = Inches(2)
        row_cells[1].width = Inches(1)
        row_cells[2].width = Inches(4)
        row_cells[3].width = Inches(2)

        row_cells[0].text = str(index)
        row_cells[1].text = "2025-04-03"
        row_cells[2].text = "ok"
        row_cells[3].text = "Yang Xie"
        row_cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row_cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row_cells[2].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row_cells[3].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    
    doc.add_page_break()
    return heading


# 2 章下列表项的项目符号编号定义（照抄参考文档 2.0.0.31 的 numId=13 → abstractNumId=9）：
# Wingdings \uF0D8（➤），缩进 left=420 hanging=420。python-docx 默认模板的编号 id 用到 0-8/1-9，
# 因此这里取 10，避免冲突。
BULLET_ABSTRACT_NUM_ID = 10
BULLET_NUM_ID = 10

_BULLET_ABSTRACT_NUM_XML = (
    '<w:abstractNum %s w:abstractNumId="%d">'
    '<w:nsid w:val="263B6CCF"/>'
    '<w:multiLevelType w:val="hybridMultilevel"/>'
    '<w:tmpl w:val="28BAB816"/>'
    '<w:lvl w:ilvl="0" w:tplc="0409000B">'
    '<w:start w:val="1"/>'
    '<w:numFmt w:val="bullet"/>'
    '<w:lvlText w:val="\uF0D8"/>'
    '<w:lvlJc w:val="left"/>'
    '<w:pPr><w:ind w:left="420" w:hanging="420"/></w:pPr>'
    '<w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr>'
    '</w:lvl>'
    '</w:abstractNum>'
)


def ensure_bullet_numbering(doc):
    """往文档的 numbering.xml 注入 ➤ 项目符号定义（只注入一次）。"""
    numbering = doc.part.numbering_part.element
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(BULLET_NUM_ID):
            return
    abstract = parse_xml(_BULLET_ABSTRACT_NUM_XML % (nsdecls("w"), BULLET_ABSTRACT_NUM_ID))
    num = parse_xml(
        '<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/></w:num>'
        % (nsdecls("w"), BULLET_NUM_ID, BULLET_ABSTRACT_NUM_ID)
    )
    # schema 要求所有 abstractNum 排在所有 num 之前，所以插到第一个 num 之前
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
        abstract.addnext(num)
    else:
        numbering.append(abstract)
        numbering.append(num)


def add_ordered_list_item(doc, text):
    """
    2 章下的列表项。参考文档 2.0.0.31 里这些条目带 ➤ 项目符号
    （pStyle=List Paragraph + numPr，符号是 Wingdings \\uF0D8），原先只是普通段落，没符号。
    """
    ensure_bullet_numbering(doc)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    paragraph.style = doc.styles["List Paragraph"]
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = BULLET_NUM_ID

    run = paragraph.add_run(text)
    run.font.size = Pt(9)  # 重置字体大小（如果需要）
    run.font.name = "Malgun Gothic"



def write_sdk_description(doc):
    add_toc_heading(doc, "2. What is the Fingerprint SDK (ISO)", 1, page=4)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("    Fingerprint SDK (ISO)")
    run.bold = True
    run.font.size = Pt(9)  # 重置字体大小（如果需要）
    run.font.name = "Malgun Gothic"
    run = paragraph.add_run(" is a fingerprint authentication software development tool which is designed for the application developers to apply the fingerprint recognition functions in developing of fingerprint recognition programs based on fingerprint international standard.")
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(9) 

    add_toc_heading(doc, "Supported Fingerprint Sensor", 2, page=4)
    add_ordered_list_item(doc, "Optical fingerprint sensor(G4 G6 G7)")

    add_toc_heading(doc, "SDK Features", 2, page=4)
    add_ordered_list_item(doc, "Includes FBI-certified, high-performance WSQ fingerprint image compression")
    add_ordered_list_item(doc, "APIs to develop applications to read, write, edit and view ISO/NIST-compliant biometric data")
    add_ordered_list_item(doc, "Run-time applications and demos with source code")
    add_ordered_list_item(doc, "Single API to create both types of structured data (binary and XML)")
    add_ordered_list_item(doc, "C++, .NET, and JNI versions available")
    add_ordered_list_item(doc, "Fully featured C Language API")
    add_ordered_list_item(doc, "C#/.NET wrappers")
    add_ordered_list_item(doc, "Example programs with source")
    add_ordered_list_item(doc, "Support for – Java Native Interface – Microsoft Windows – Linux – and Android")
  
    add_toc_heading(doc, "Support for", 2, page=4)
    add_ordered_list_item(doc, "Java Native Interface")
    add_ordered_list_item(doc, "Microsoft Windows")
    add_ordered_list_item(doc, "Linux")
    add_ordered_list_item(doc, "Android")

    add_toc_heading(doc, "Supported Communication Interface", 2, page=4)
    add_ordered_list_item(doc, "USB 1.0, 1.1 and 2.0 (Full Speed) specifications")
    doc.add_page_break()


def createSDKAPI(config, verName, APIPathName):

    liner = readAPI(APIPathName, config["keyword"], config["funKey"])
    if len(liner) < 1:
        print("liner is null keyword {keyword}".format(keyword= config["keyword"] ))
        return

    _TOC_ENTRIES.clear()
    _BOOKMARK_ID[0] = 1000

    doc = Document()
    # doc.add_paragraph('hello world')
    heading = doc.add_heading('Fingerprint Scanner SDK (ISO)', 0)
    paragraph = heading._element
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    heading = doc.add_heading('User Manual', 0)
    paragraph = heading._element
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    heading = doc.add_heading('V'+verName, 0)
    paragraph = heading._element
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for index in range(17):
        doc.add_paragraph("")

    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Shanghai OXi Technology Company Ltd.")
   
    run.font.size = Pt(13)  # 重置字体大小（如果需要）
    run.font.name = "Malgun Gothic"
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_page_break()

    rev_heading = write_revision_record(doc)

    write_sdk_description(doc)

    # API References
    add_toc_heading(doc, "3. API References", 1, page=5)
    doc.add_paragraph("The library-calling interface includes the following functions: ")

    table = doc.add_table(rows=1, cols=3, style="Table Grid")
    hdr_cells = table.rows[0].cells
    table.rows[0].height = Inches(0.3)
    hdr_cells[0].text = 'No'
    hdr_cells[1].text = 'function'
    hdr_cells[2].text = 'explain'

    for cell in table.rows[0].cells:
        set_cell_background(cell, 'c0c0c0')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for index in range(len(liner)):
        row_cells = table.add_row().cells
        table.rows[index+1].height = Inches(0.3)
        row_cells[0].text = str(index)
        row_cells[0].width = Inches(2)
        row_cells[1].width = Inches(4)
        row_cells[2].width = Inches(2)
        row_cells[1].text = extract_function_name(liner[index])
        row_cells[2].text = "ok"
        row_cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row_cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row_cells[2].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

   #document.add_paragraph('first item in unordered list', style='List Bullet')
    for index in range(len(liner)):
        funName = extract_function_name(liner[index])
        print(funName)
        add_toc_heading(doc, str(index+1) + "." + funName, 2, page=5 + index // 2)

        #### Signature
        sin = doc.add_paragraph()
        run = sin.add_run("Signature")
        run.font.bold = True
        run.font.italic = True
       
        table2 = doc.add_table(rows=1, cols=2,style="Table Grid")

        table2.rows[0].height = Inches(0.4)
        hdr_cells = table2.rows[0].cells
        hdr_cells[0].width = Inches(1.2)
        set_cell_background(hdr_cells[0], 'c9c9c9')
        hdr_cells[0].text = "Method"
        hdr_cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        hdr_cells[0].paragraphs[0].runs[0].bold = True
        hdr_cells[1].text = getFuntionString(liner[index], config["keyword"])
        hdr_cells[1].width = Inches(8)
        hdr_cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        
        funInfo = extract_brief_description(funName, APIPathName)

        sin = doc.add_paragraph()
        sin.add_run("\n")
        run = sin.add_run("Description")
        run.font.bold = True
        run.font.italic = True
        sin.add_run("\n")
       
        for desc in funInfo:
            if "brief" in desc:
                run = sin.add_run("      " + desc.split("brief")[1].strip() + "\n")
                break

        res = doc.add_paragraph()
        run = res.add_run("Result")
        res.add_run("\n")
        run.font.bold = True
        run.font.italic = True
        havaReturn = False
        for desc in funInfo:
            if "return" in desc:
                havaReturn = True
                continue
            if havaReturn:
                if config["keyword"] not in desc:
                     res.add_run(desc.replace("///", "").replace("\t","    "))
                     res.add_run("\n")

        sin = doc.add_paragraph()
        run = sin.add_run("Parameter")
        run.font.bold = True
        run.font.italic = True
        para = []
        for desc in funInfo:
            if "param" in desc:
                info = desc.split("param")[1].strip()
                para.append(info)
                print(info)
        
        tableLen = len(para) + 1
        if(tableLen > 1):
            table3 = doc.add_table(rows=1, cols=2,style="Table Grid")
            table3.rows[0].height = Inches(0.3)
            hdr_cells = table3.rows[0].cells
            hdr_cells[0].text = "Parameter"
            hdr_cells[0].width = Inches(2)
            hdr_cells[0].paragraphs[0].runs[0].bold = True
            hdr_cells[1].text = "Description"
            hdr_cells[1].width = Inches(6)
            hdr_cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            hdr_cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_background(hdr_cells[0], 'c9c9c9')
            set_cell_background(hdr_cells[1], 'c9c9c9')

            for index in range(len(para)):
                cs = para[index].split("[")
                row_cells = table3.add_row().cells
                table3.rows[index+1].height = Inches(0.3)
                row_cells[0].text = cs[0]
                row_cells[0].width = Inches(2)
                row_cells[1].width = Inches(6)
                row_cells[1].text = para[index].replace(cs[0], "").strip()
                row_cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                row_cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        sin = doc.add_paragraph()
        sin.add_run("\n")

    # 正文写完，回填目录页（封面之后、第 1 章之前）
    insert_toc(doc, rev_heading._p, _TOC_ENTRIES)

    doc.save(config["output_name"].format(version=verName))
    

 # 添加有序列表项的函数
    
   

def load_config():
    with open("api_config.json", encoding="utf-8") as f:
        return json.load(f)


if __name__ == '__main__':

    config = load_config()

     # config["GSER","A910", "MIAXIS", "huantian", "Thales",]
    sdk = config["GSER"]
    if not sdk["sdk_dir"] == "":
        print(sdk["sdk_dir"])
        result = parse_cmake_versions(sdk["sdk_dir"])
        print(sdk["output_name"].format(version=result['fap_version']))

        print(f"--------------{result['fap_version']}---------------")
        APIFilePath = findFilePathNameByKeyWord(sdk["sdk_dir"], sdk["macro"]) 
        print(APIFilePath)
        if APIFilePath is None:
            print("APIFilePath is null")
        else:
            createSDKAPI(sdk, result['fap_version'], APIFilePath)
    # createSDKAPI("D:\\git\\oxi\\OXiScan\\dev\\fap70", "D:\\G7SDKLib\\doc")
    
   